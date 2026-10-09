from datetime import datetime, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from basics.models import Balance, Currency, PaymentSystem, TrafficType
from merchant.models import Merchant, MerchantSolution
from trade.models import InOrder, InOrderStatus, WithdrawalRequest


class MerchantLkExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="aggrepay", password="x")
        usdt = Balance.objects.create(type=0, amount=Decimal("10000"))
        frozen = Balance.objects.create(type=1, amount=Decimal("0"))
        self.merchant = Merchant.objects.create(user=self.user, balance=usdt, frozen_balance=frozen)

        kzt = Currency.objects.create(name="Tenge LK", symbol="KZTLK")
        ps = PaymentSystem.objects.create(
            name="C2CKZT",
            currency=kzt,
            required_fields={},
            usdt_exchange_rate=Decimal("500"),
        )
        traffic = TrafficType.objects.create(name="Standard LK")
        self.solution = MerchantSolution.objects.create(
            merchant=self.merchant,
            payment_system=ps,
            traffic=traffic,
            mdr_in=Decimal("7.50"),
        )
        self.completed = InOrderStatus.objects.create(name="Completed")
        self.order = InOrder.objects.create(
            status=self.completed,
            amount=Decimal("4000.00"),
            usd_amount=Decimal("8.00"),
            solution=self.solution,
            merchant_order_id="agg-sep-1",
            merchant_fee=Decimal("300.00"),
            trader_fee=Decimal("0.40"),
            arbitrage_comment="",
            pic="https://example.com/receipt",
        )
        InOrder.objects.filter(pk=self.order.pk).update(
            creation_date=timezone.make_aware(datetime(2026, 9, 15, 12, 0, 0), dt_timezone.utc),
        )
        self.order.refresh_from_db()

        self.withdrawal = WithdrawalRequest.objects.create(
            status=1,
            amount=Decimal("4000.00"),
            balance=usdt,
            address_to="TWtestaddress000000000000000000000",
            comment="settlement",
            from_user=self.user,
        )
        WithdrawalRequest.objects.filter(pk=self.withdrawal.pk).update(
            date=timezone.make_aware(datetime(2026, 9, 20, 10, 0, 0), dt_timezone.utc),
        )
        self.withdrawal.refresh_from_db()

        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_orders_in_rejects_broken_flatpickr_range(self):
        response = self.client.get("/api/v1/trade/order/in/", {
            "creation_date__range": "2026-08-31 12:00,2026-10-01 12:00T23:59:59",
            "status__name__in": "Completed",
        })
        self.assertEqual(response.status_code, 400)
        self.assertIn("creation_date__range", response.data)

    def test_orders_in_accepts_iso_datetime_range(self):
        response = self.client.get("/api/v1/trade/order/in/", {
            "creation_date__range": "2026-09-01T00:00:00,2026-09-30T23:59:59",
            "status__name__in": "Completed",
        })
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertIn(str(self.order.id), ids)

    def test_orders_in_payment_system_alias_filters(self):
        response = self.client.get("/api/v1/trade/order/in/", {
            "creation_date__range": "2026-09-01T00:00:00,2026-09-30T23:59:59",
            "status__name__in": "Completed",
            "payment_system__name__in": "C2CKZT",
        })
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertIn(str(self.order.id), ids)

        empty = self.client.get("/api/v1/trade/order/in/", {
            "creation_date__range": "2026-09-01T00:00:00,2026-09-30T23:59:59",
            "status__name__in": "Completed",
            "payment_system__name__in": "MissingPS",
        })
        self.assertEqual(empty.status_code, 200)
        self.assertEqual(empty.data["results"], [])

    def test_orders_in_export_uses_same_iso_range(self):
        response = self.client.get("/api/v1/trade/order/in/export/", {
            "creation_date__range": "2026-09-01T00:00:00,2026-09-30T23:59:59",
            "status__name__in": "Completed",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        values = [str(cell.value) for row in sheet.iter_rows(min_row=2) for cell in row]
        self.assertTrue(any(str(self.order.id) in value for value in values))

    def test_withdrawal_export_returns_xlsx(self):
        response = self.client.get("/api/v1/trade/withdrawal-request/export/", {
            "date__range": "2026-09-01T00:00:00,2026-09-30T23:59:59",
            "status__in": "1",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        self.assertIn("ID", headers)
        self.assertIn("Статус", headers)
        self.assertIn("Сумма (USDT)", headers)
        self.assertIn("Адрес", headers)
        id_idx = headers.index("ID")
        status_idx = headers.index("Статус")
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 1)
        self.assertEqual(str(rows[0][id_idx]), str(self.withdrawal.id))
        self.assertEqual(rows[0][status_idx], "Success")
