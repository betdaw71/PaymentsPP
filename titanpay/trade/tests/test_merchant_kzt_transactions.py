from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from basics.models import Balance, Currency, PaymentSystem, TrafficType
from merchant.kzt_settlement import ensure_kzt_balances
from merchant.models import Merchant, MerchantSolution
from trade.ledger import merchant_balance_ids
from trade.models import InOrder, InOrderStatus, OutOrder, OutOrderStatus, Transaction, TransactionType


class MerchantKztTransactionsApiTest(TestCase):
    def setUp(self):
        self.deposit_type, _ = TransactionType.objects.get_or_create(name="Deposit")
        self.charge_type, _ = TransactionType.objects.get_or_create(name="Charge")
        self.blockchain, _ = Balance.objects.get_or_create(type=3, defaults={"amount": Decimal("1000000")})
        self.blockchain.amount = Decimal("1000000")
        self.blockchain.save(update_fields=["amount"])

        self.user = User.objects.create_user(username="melbet_kzt_tx", password="x")
        usdt = Balance.objects.create(type=0, amount=Decimal("0"))
        frozen = Balance.objects.create(type=1, amount=Decimal("0"))
        self.merchant = Merchant.objects.create(user=self.user, balance=usdt, frozen_balance=frozen)
        ensure_kzt_balances(self.merchant)
        self.merchant.refresh_from_db()

        kzt = Currency.objects.create(name="Tenge TX", symbol="KZTX")
        ps = PaymentSystem.objects.create(
            name="C2CKZT",
            currency=kzt,
            required_fields={},
            usdt_exchange_rate=Decimal("500"),
        )
        traffic = TrafficType.objects.create(name="Standard TX")
        solution = MerchantSolution.objects.create(
            merchant=self.merchant,
            payment_system=ps,
            traffic=traffic,
            mdr_in=Decimal("7.50"),
        )
        status = InOrderStatus.objects.create(name="Completed")
        self.order = InOrder.objects.create(
            status=status,
            amount=Decimal("5000.00"),
            usd_amount=Decimal("10.00"),
            solution=solution,
            merchant_order_id="melbet-kzt-1",
            merchant_fee=Decimal("375.00"),
            trader_fee=Decimal("0.60"),
            arbitrage_comment="",
            pic="https://example.com/receipt",
        )
        self.kzt_tx = Transaction.create(
            self.blockchain,
            self.merchant.balance_kzt,
            value=Decimal("4625.00"),
            _transaction_type=self.deposit_type,
            _linked_in_order=self.order,
            _comment="KZT deposit",
        )
        self.usdt_tx = Transaction.create(
            self.blockchain,
            self.merchant.balance,
            value=Decimal("12.34"),
            _transaction_type=self.deposit_type,
            _comment="USDT deposit",
        )

        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_merchant_balance_ids_include_kzt(self):
        ids = set(merchant_balance_ids(self.merchant))
        self.assertIn(self.merchant.balance_id, ids)
        self.assertIn(self.merchant.balance_kzt_id, ids)
        self.assertIn(self.merchant.frozen_balance_kzt_id, ids)

    def test_list_includes_kzt_and_usdt_deposits(self):
        response = self.client.get("/api/v1/trade/transaction/", {"transaction_type__name__in": "Deposit"})
        self.assertEqual(response.status_code, 200)
        rows = {item["id"]: item for item in response.data["results"]}
        self.assertIn(str(self.kzt_tx.id), rows)
        self.assertIn(str(self.usdt_tx.id), rows)
        self.assertEqual(rows[str(self.kzt_tx.id)]["currency"], "KZT")
        self.assertTrue(rows[str(self.kzt_tx.id)]["is_incoming"])
        self.assertEqual(Decimal(str(rows[str(self.kzt_tx.id)]["fee"])), Decimal("375.00"))
        self.assertEqual(Decimal(str(rows[str(self.kzt_tx.id)]["value"])), Decimal("5000.00"))
        self.assertEqual(rows[str(self.kzt_tx.id)]["order_status"], "Completed")
        self.assertEqual(rows[str(self.usdt_tx.id)]["currency"], "USDT")

    def test_incoming_direction_includes_kzt_balance(self):
        response = self.client.get("/api/v1/trade/transaction/", {
            "direction": "incoming",
            "transaction_type__name__in": "Deposit",
        })
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertIn(str(self.kzt_tx.id), ids)
        self.assertIn(str(self.usdt_tx.id), ids)

    def test_legacy_available_merchant_filter_misses_kzt(self):
        response = self.client.get("/api/v1/trade/transaction/", {
            "to_balance__available_merchant__user__username__in": self.user.username,
            "transaction_type__name__in": "Deposit",
        })
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.data["results"]}
        self.assertIn(str(self.usdt_tx.id), ids)
        self.assertNotIn(str(self.kzt_tx.id), ids)

    def test_export_includes_amount_fee_status_and_kzt_currency(self):
        response = self.client.get("/api/v1/trade/transaction/export/", {
            "direction": "incoming",
            "transaction_type__name__in": "Deposit",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            response["Content-Type"],
        )
        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        self.assertIn("ID (транзакции)", headers)
        self.assertIn("Сумма", headers)
        self.assertIn("Комиссия мерчанта", headers)
        self.assertIn("Статус заявки", headers)
        self.assertIn("Дата", headers)
        self.assertIn("Валюта", headers)

        id_idx = headers.index("ID (транзакции)")
        currency_idx = headers.index("Валюта")
        fee_idx = headers.index("Комиссия мерчанта")
        status_idx = headers.index("Статус заявки")
        amount_idx = headers.index("Сумма")
        rows = {
            str(row[id_idx].value): row
            for row in sheet.iter_rows(min_row=2)
        }
        kzt_row = rows[str(self.kzt_tx.id)]
        self.assertEqual(kzt_row[currency_idx].value, "KZT")
        self.assertEqual(Decimal(str(kzt_row[fee_idx].value)), Decimal("375.00"))
        self.assertEqual(Decimal(str(kzt_row[amount_idx].value)), Decimal("5000.00"))
        self.assertEqual(kzt_row[status_idx].value, "Completed")

    def test_withdrawal_filter_includes_payout_charge(self):
        out_status = OutOrderStatus.objects.create(name="Completed")
        out_order = OutOrder.objects.create(
            status=out_status,
            amount=Decimal("20000.00"),
            usd_amount=Decimal("40.00"),
            solution=self.order.solution,
            merchant_order_id="melbet-kzt-out-1",
            merchant_fee=Decimal("1500.00"),
            trader_fee=Decimal("1.00"),
            pic="https://example.com/out",
        )
        payout_tx = Transaction.create(
            self.merchant.frozen_balance_kzt,
            self.blockchain,
            value=Decimal("21500.00"),
            _transaction_type=self.charge_type,
            _linked_out_order=out_order,
            _comment="Out-order completed",
        )
        response = self.client.get("/api/v1/trade/transaction/", {
            "direction": "outgoing",
            "transaction_type__name__in": "Withdrawal",
        })
        self.assertEqual(response.status_code, 200)
        rows = {item["id"]: item for item in response.data["results"]}
        self.assertIn(str(payout_tx.id), rows)
        self.assertEqual(rows[str(payout_tx.id)]["transaction_type"], "Withdrawal")
        self.assertEqual(Decimal(str(rows[str(payout_tx.id)]["value"])), Decimal("20000.00"))
        self.assertEqual(Decimal(str(rows[str(payout_tx.id)]["fee"])), Decimal("1500.00"))
        self.assertNotIn(str(self.kzt_tx.id), rows)
