#!/usr/bin/env python3
"""Fill Melbet KG (KGS) AvaPay integration questionnaires with credentials and commercial terms."""

from __future__ import annotations

from pathlib import Path

from docx import Document

DIR = Path(__file__).resolve().parent

API_KEY = "vJ-57rbr1kVRsUNo9xKiGNOqtBFal-G3gnpC7RZCNwo"
SECRET = "2d26e0fd362931148aef4e71e3f93f5d595d4300539949b916045384d01b405b"

CREDS = (
    f"x-api-key: {API_KEY}\n"
    f"Secret key (HMAC-SHA256 для x-signature): {SECRET}\n"
    "POST deposit (RedirectURL): https://api.avapay.net/api/v1/integrations/melbet/deposit/\n"
    "POST withdrawal (host-to-host): https://api.avapay.net/api/v1/integrations/melbet/withdrawal/\n"
    "GET status: https://api.avapay.net/api/v1/integrations/melbet/transactions/{transaction_id}/\n"
    "Те же ключи на все KG-методы: qr_kgs, phone_kgs, card2card_kgs (PayPlat / AvaPay)."
)

ADMIN_PANEL = (
    "https://avapay.net/ — логин кабинета мерчанта уже выдан. "
    "Пароль сменён вашим менеджером; актуального пароля у нас нет, доступа к нему не имеем. "
    "По вопросам доступа / сброса — Telegram: @avapay_support5"
)

SUPPORT = "Telegram: @avapay_support5 (операционные инциденты и доступ в ЛК)."

USAGE = (
    "Сайты: Melbet (betting), по договору. GEO: Киргизия (KG). "
    "Валюта запроса / процессинга: KGS. Валюта баланса и settlement: USDT. "
    "Тип трафика: первичный и вторичный. Заявленная конверсия: 50%. "
    "Тип интеграции: H2H / Redirect. Провайдер: PayPlat (AvaPay). "
    "Лимиты по MID (месячные/суточные): неограниченно. "
    "Settlement: USDT, fee 0%, T+0, от 3 000 USDT. Курс: XE (https://www.xe.com) + 2%."
)

FEE_IN = "7% (валюта комиссии — USDT; amount в API = GROSS к оплате клиентом в KGS)."
FEE_OUT = "3.2% (валюта комиссии — USDT)."
LIMITS_IN = "KGS: min 500 / max 1 000 000."
LIMITS_OUT = (
    "KGS: min 5 000 / max 1 000 000. "
    "Out info: выплаты по QR, 1–3 чека в зависимости от величины суммы."
)

DEPOSIT_NOTIFY = (
    "POST на callback_url из заявки.\n"
    "Headers:\n"
    "  Content-Type: application/json\n"
    "  x-api-key: <ваш x-api-key>\n"
    "  x-signature: <HMAC-SHA256 hex от raw body>\n"
    "Body (пример):\n"
    '{"order_id":"MB-123456789"}\n'
    "Суммы в callback нет. После callback — GET /transactions/{transaction_id}/ "
    "→ status, amount (GROSS KGS), currency, method."
)

WITHDRAW_NOTIFY = (
    "Тот же формат, что депозит.\n"
    "POST на callback_url.\n"
    "Headers: Content-Type: application/json; x-api-key; x-signature (HMAC-SHA256 raw body).\n"
    "Body (пример):\n"
    '{"order_id":"MB-W-987654321"}\n'
    "Далее GET /transactions/{transaction_id}/ → status, amount (GROSS KGS)."
)

FX_TABLE3_DIFF = (
    "Да. Валюта процессинга / запроса: KGS. Валюта баланса и settlement: USDT. "
    "Источник курса: https://www.xe.com (XE) + курсовая маржа 2%. "
    "Settlement: USDT, fee 0%, период T+0, лимит от 3 000 USDT. "
    "В Melbet API суммы передаются в KGS (валюта запроса); конвертация в USDT — на стороне settlement."
)

FX_GROSS_IN = (
    "В API Melbet отдельного поля суммы в USDT нет. "
    "GROSS депозита в валюте запроса: параметр amount (KGS) в POST /deposit/ и GET /transactions/{id}/. "
    "Эквивалент в валюте баланса (USDT) считается при settlement по курсу XE + 2%. "
    "В callback суммы нет — только {\"order_id\":\"...\"}."
)

FX_GROSS_OUT = (
    "В API Melbet отдельного поля суммы в USDT нет. "
    "GROSS вывода в валюте запроса: параметр amount (KGS) в POST /withdrawal/ и GET /transactions/{id}/. "
    "Эквивалент в USDT — при settlement (XE + 2%). В callback суммы нет."
)

VIDEO_NOTE = "Видео флоу оплаты — предоставим отдельным файлом / по запросу."


def set_cell(cell, text: str) -> None:
    if not cell.paragraphs:
        cell.text = text
        return
    cell.paragraphs[0].text = text
    for p in cell.paragraphs[1:]:
        p.text = ""


def comment_cells(row):
    """Return comment cell indices that are unique in the row (handles merged duplicates)."""
    seen = {}
    for i, c in enumerate(row.cells):
        seen.setdefault(id(c._tc), i)
    # Last unique cells are usually Comment columns; for table0 it's cells[3], for table1 cells[4]/5]
    return list(seen.values())


def set_comment(row, text: str, prefer_last_n: int = 1) -> None:
    idxs = comment_cells(row)
    # Prefer rightmost unique cells (Comment columns)
    targets = idxs[-prefer_last_n:] if prefer_last_n else idxs[-1:]
    for i in targets:
        set_cell(row.cells[i], text)


def fill_common_provider(doc: Document, method: str, method_label: str) -> None:
    t0 = doc.tables[0]
    set_comment(
        t0.rows[1],
        "Типовая интеграция Melbet (спецификация v3.1).\n"
        "База: https://api.avapay.net/api/v1/integrations/melbet/\n"
        "POST deposit/  POST withdrawal/  GET transactions/{transaction_id}/\n"
        "Публичная merchant-дока: https://doc.avapay.net/\n"
        f"Видео флоу: {VIDEO_NOTE}",
    )
    set_comment(t0.rows[2], "AvaPay (PayPlat)")
    set_comment(
        t0.rows[3],
        f"KG / KGS / {method_label} / method={method}\n"
        "Трафик: первичный + вторичный. Конверсия: 50%. "
        "Интеграция: H2H/Redirect. Баланс: USDT, запрос: KGS.\n"
        "Все KG-методы одной интеграции PayPlat: qr_kgs (ELQR), phone_kgs (мобильный), "
        "card2card_kgs (карта). Магазин и ключи общие.",
    )
    set_comment(t0.rows[4], CREDS)
    set_comment(
        t0.rows[5],
        "Да. По запросу мерчанта ротируем секретный ключ в бэк-офисе, старый деактивируется, "
        "выдаём новую пару. Контакт: @avapay_support5",
    )
    set_comment(t0.rows[6], ADMIN_PANEL)
    set_comment(
        t0.rows[7],
        "Доступ в ЛК и API можно ограничить whitelist IP. 2FA кабинета — по настройке аккаунта / по запросу. "
        "Пароль ЛК сменён менеджером Melbet — у нас нет доступа; поддержка: @avapay_support5",
    )
    set_comment(t0.rows[8], USAGE)
    set_comment(
        t0.rows[9],
        "Тестовый контур и тестовые ключи — по запросу. Боевые ключи (выше) с тестом не смешиваем.",
    )
    set_comment(
        t0.rows[10],
        "AvaPay / PayPlat. Операционный контакт поддержки: Telegram @avapay_support5. "
        "Менеджер по текущей переписке — в сопроводительном письме.",
    )
    set_comment(t0.rows[11], SUPPORT)
    set_comment(
        t0.rows[12],
        "Исходящие callback с серверов API AvaPay. Актуальный список IP вышлем отдельным письмом / по запросу.",
    )
    set_comment(
        t0.rows[13],
        "Да. Просим прислать ваши исходящие IP для whitelist на приём deposit / withdrawal / status.",
    )
    set_comment(
        t0.rows[14],
        "Да. order_id из вашего запроса. Возвращаем его в callback и в GET status.",
    )
    set_comment(
        t0.rows[15],
        "Да. Повторный order_id того же мерчанта → HTTP 409, заказ уже существует.",
    )
    set_comment(t0.rows[16], "SUCCESS (внутренний Success).")
    set_comment(t0.rows[17], "FAILED (внутренние Failed / Declined / Expired).")
    set_comment(t0.rows[18], "SUCCESS (внутренний Success). Выводы доступны (PayPlat / AvaPay).")
    set_comment(t0.rows[19], "FAILED (внутренние Failed / Declined / Expired).")
    set_comment(t0.rows[20], DEPOSIT_NOTIFY)
    set_comment(t0.rows[21], WITHDRAW_NOTIFY)
    set_comment(
        t0.rows[24],
        "Ошибки: HTTP 400 / 401 / 409, тело {\"error\":{\"code\":<int>,\"message\":\"<string>\"}}.\n"
        "401 — нет или неверные x-api-key / x-signature, либо IP не в whitelist.\n"
        "409 — дубль order_id или у клиента уже есть pending.\n"
        "400 — валидация (нет order_id / amount / currency / method, метод не активен, лимиты, нет реквизитов и т.д.).",
    )
    set_comment(
        t0.rows[25],
        "Да. amount — число, до 2 знаков после запятой (пример: 1500.50).",
    )


def fill_payment_methods_common(doc: Document, method: str, button: str, deposit_flow: str, withdraw_extra: str) -> None:
    t1 = doc.tables[1]
    # prefer_last_n=2 for duplicated comment columns
    set_comment(
        t1.rows[2],
        "Депозит: Redirect (redirect_url). Вывод: H2H (host-to-host, без редиректа). Не iframe.",
        prefer_last_n=2,
    )
    set_comment(
        t1.rows[3],
        f"На стороне мерчанта (передаёте method={method}).",
        prefer_last_n=2,
    )
    set_comment(t1.rows[4], "Депозит и вывод.", prefer_last_n=2)
    set_comment(t1.rows[5], "KGS", prefer_last_n=2)
    set_comment(t1.rows[6], "KGS", prefer_last_n=2)
    set_comment(t1.rows[7], FEE_IN, prefer_last_n=2)
    set_comment(t1.rows[8], "Нет (фиксированной комиссии нет).", prefer_last_n=2)
    set_comment(t1.rows[9], LIMITS_IN, prefer_last_n=2)
    set_comment(t1.rows[10], "KGS", prefer_last_n=2)
    set_comment(t1.rows[11], FEE_OUT, prefer_last_n=2)
    set_comment(t1.rows[12], "Нет (фиксированной комиссии нет).", prefer_last_n=2)
    set_comment(t1.rows[13], LIMITS_OUT, prefer_last_n=2)
    set_comment(t1.rows[14], button, prefer_last_n=2)
    set_comment(
        t1.rows[15],
        "Предоставим по запросу (150×50 px или соотношение 3:1).",
        prefer_last_n=2,
    )
    set_comment(
        t1.rows[19],
        "Да: лимиты, pending, blacklist, уникальный order_id, IP whitelist.",
        prefer_last_n=2,
    )
    set_comment(
        t1.rows[23],
        "Нет. HMAC-SHA256 (заголовок x-signature) + x-api-key.",
        prefer_last_n=1,
    )
    set_comment(
        t1.rows[24],
        f"order_id, amount, currency=kgs, method={method}, callback_url, "
        "success_url, pending_url, fail_url, customer_account_id. "
        "Желательно fields.email / phone / first_name / last_name.",
        prefer_last_n=1,
    )
    set_comment(
        t1.rows[25],
        f"order_id, amount, currency=kgs, method={method}, callback_url, "
        f"customer_account_id, account_number ({withdraw_extra}).",
        prefer_last_n=1,
    )
    set_comment(
        t1.rows[26],
        "fields.first_name, last_name, phone, email.",
        prefer_last_n=1,
    )
    set_comment(
        t1.rows[27],
        "fields.first_name, last_name, email.",
        prefer_last_n=1,
    )

    t0 = doc.tables[0]
    set_comment(t0.rows[26], deposit_flow)
    set_comment(
        t0.rows[27],
        "Валюта процессинга (KGS) отличается от валюты баланса (USDT). "
        "В Melbet API суммы — в KGS: amount в POST и GET status, currency=kgs. "
        "Отдельного поля USDT в callback/status нет; USDT — settlement по XE+2%. "
        "Callback: только {\"order_id\":\"...\"}.",
    )


def fill_fx(doc: Document) -> None:
    t3 = doc.tables[3]
    set_comment(t3.rows[0], FX_TABLE3_DIFF)
    set_comment(t3.rows[1], FX_GROSS_IN)
    set_comment(t3.rows[2], FX_GROSS_OUT)


def fill_chargeback(doc: Document) -> None:
    t2 = doc.tables[2]
    set_comment(
        t2.rows[0],
        "Не карточный эквайринг Visa/MC: классических chargeback по схеме нет. "
        "Спор по P2P — апелляция / ручной разбор. Статус — callback {\"order_id\"} + GET status. "
        "Файлы и скрины — Telegram @avapay_support5 / тикет.",
    )
    set_comment(
        t2.rows[1],
        "Лимиты суммы, один pending pay-in на клиента, blacklist клиента, уникальный order_id, "
        "опционально IP whitelist API, срок жизни заявки.",
    )
    set_comment(
        t2.rows[2],
        "Вручную после разбора. Автовозврата через API нет. "
        "Итог: финальный FAILED либо пересчёт amount при подтверждённой иной фактической оплате.",
    )
    set_comment(
        t2.rows[3],
        "1) Тикет / апелляция (@avapay_support5). 2) Проверка платежа. "
        "3) Решение: SUCCESS / FAILED / пересчёт amount. 4) Callback + актуальный GET status.",
    )


def fill_qr(path: Path) -> None:
    doc = Document(path)
    method = "qr_kgs"
    if doc.paragraphs:
        doc.paragraphs[0].text = (
            "ЗАПОЛНЕНО: AvaPay/PayPlat + qr_kgs + KG | Тема: AvaPay + qr_kgs + KG | "
            "Креды, комиссии, лимиты и выводы внесены. Поддержка: @avapay_support5"
        )
    fill_common_provider(doc, method, "ELQR / QR")
    set_comment(
        doc.tables[0].rows[22],
        "Не применимо к qr_kgs (не мобильный метод). Мобильный метод KG — отдельная анкета phone_kgs "
        "(тот же магазин/ключи PayPlat).",
    )
    set_comment(
        doc.tables[0].rows[23],
        "Не обязателен для QR. Если передаёте fields.phone — E.164, KG: +996XXXXXXXXX без пробелов.",
    )
    fill_payment_methods_common(
        doc,
        method,
        "QR KGS / ELQR",
        "Redirect: открываете redirect_url из ответа депозита (страница/виджет с QR).\n"
        "На странице: сумма KGS, QR для оплаты (ELQR), таймер жизни заявки.\n"
        "Клиент сканирует QR и оплачивает.\n"
        "Вывод: H2H POST /withdrawal/ method=qr_kgs; выплаты по QR — 1–3 чека в зависимости от суммы.\n"
        f"{VIDEO_NOTE}",
        "реквизиты/идентификатор выплаты по QR (по договорённости канала)",
    )
    set_comment(
        doc.tables[1].rows[16],
        "Депозит: нет (мы выдаём QR). Вывод: реквизиты получателя передаёте в запросе вывода.",
        prefer_last_n=2,
    )
    set_comment(doc.tables[1].rows[17], "Нет (не карточный метод).", prefer_last_n=2)
    set_comment(doc.tables[1].rows[18], "Н/д (не карточный метод).", prefer_last_n=2)
    set_comment(doc.tables[1].rows[20], "Нет.", prefer_last_n=2)
    set_comment(
        doc.tables[1].rows[21],
        "QR-код (ELQR) на странице оплаты, сумма KGS. Выплаты out — по QR, 1–3 чека.",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[0].rows[28],
        "Депозит: 1) POST /deposit/ (amount, currency=kgs, method=qr_kgs, order_id, callback_url, "
        "success/pending/fail_url, customer). 2) 201: transaction_id, redirect_url. "
        "3) Клиент оплачивает по QR. 4) POST callback {\"order_id\":\"...\"}. "
        "5) GET /transactions/{id}/ → status, amount (GROSS KGS).\n"
        "Вывод: POST /withdrawal/ method=qr_kgs (H2H) → transaction_id → callback {\"order_id\"} → GET status. "
        "Out: 1–3 QR-чека в зависимости от суммы.",
    )
    fill_chargeback(doc)
    set_comment(doc.tables[2].rows[4], "Нет (не мобильный метод).")
    set_comment(doc.tables[2].rows[5], "Нет (не карточный метод).")
    fill_fx(doc)
    doc.save(path)


def fill_phone(path: Path) -> None:
    doc = Document(path)
    method = "phone_kgs"
    if doc.paragraphs:
        doc.paragraphs[0].text = (
            "ЗАПОЛНЕНО: AvaPay/PayPlat + phone_kgs + KG (мобильный метод) | "
            "Тема: AvaPay + phone_kgs + KG | Креды, комиссии, лимиты внесены. "
            "Поддержка: @avapay_support5"
        )
    fill_common_provider(doc, method, "перевод по номеру телефона (мобильный)")
    set_comment(
        doc.tables[0].rows[22],
        "KG, код страны 996. Операторы: MegaCom, Beeline, O!, Nurtelecom и др. "
        "Номер в E.164 без пробелов. Метод phone_kgs входит в ту же интеграцию PayPlat "
        "вместе с qr_kgs и card2card_kgs.",
    )
    set_comment(
        doc.tables[0].rows[23],
        "E.164: +996XXXXXXXXX (плюс, код страны, 9 цифр абонента, без пробелов и скобок). "
        "Пример: +996700123456. Для вывода — account_number в том же формате.",
    )
    fill_payment_methods_common(
        doc,
        method,
        "Phone KGS / Перевод на телефон (KG)",
        "Депозит: Redirect, redirect_url. На странице: сумма KGS, номер телефона получателя / "
        "инструкция перевода, копирование, таймер.\n"
        "Вывод: H2H — в POST /withdrawal/ передаёте account_number = телефон получателя +996…\n"
        f"{VIDEO_NOTE}",
        "телефон получателя +996XXXXXXXXX",
    )
    set_comment(
        doc.tables[1].rows[16],
        "Депозит: нет (мы выдаём реквизиты). Вывод: да — номер телефона получателя в account_number.",
        prefer_last_n=2,
    )
    set_comment(doc.tables[1].rows[17], "Нет (не карточный метод).", prefer_last_n=2)
    set_comment(doc.tables[1].rows[18], "Н/д (не карточный метод).", prefer_last_n=2)
    set_comment(
        doc.tables[1].rows[20],
        "Нет автоматической проверки принадлежности номера. Вывод идёт на номер из запроса.",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[1].rows[21],
        "Номер телефона KG в формате +996XXXXXXXXX.",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[0].rows[28],
        "Депозит: POST /deposit/ method=phone_kgs → redirect_url → клиент переводит по номеру → "
        "callback {\"order_id\":\"...\"} → GET status (amount GROSS KGS).\n"
        "Вывод: POST /withdrawal/ (amount, currency=kgs, method=phone_kgs, order_id, "
        "account_number=+996…, callback_url, customer) → transaction_id → "
        "callback {\"order_id\"} → GET status.",
    )
    fill_chargeback(doc)
    set_comment(
        doc.tables[2].rows[4],
        "Номер, с которого платили, автоматически не отдаём. На вывод вы сами передаёте account_number.",
    )
    set_comment(doc.tables[2].rows[5], "Нет (не карточный метод).")
    fill_fx(doc)
    doc.save(path)


def fill_card(path: Path) -> None:
    doc = Document(path)
    method = "card2card_kgs"
    if doc.paragraphs:
        doc.paragraphs[0].text = (
            "ЗАПОЛНЕНО: AvaPay/PayPlat + card2card_kgs + KG | Тема: AvaPay + card2card_kgs + KG | "
            "Креды, комиссии, лимиты и выводы внесены. Поддержка: @avapay_support5"
        )
    fill_common_provider(doc, method, "P2P card-to-card / перевод по номеру карты")
    set_comment(
        doc.tables[0].rows[22],
        "Не применимо к card2card_kgs (не мобильный метод). Мобильный метод — phone_kgs "
        "(та же интеграция PayPlat / те же ключи).",
    )
    set_comment(
        doc.tables[0].rows[23],
        "Не обязателен. Если передаёте fields.phone — E.164, для KG: +996XXXXXXXXX без пробелов.",
    )
    fill_payment_methods_common(
        doc,
        method,
        "Card2Card KGS / Перевод на карту (KG)",
        "Redirect: открываете redirect_url из ответа депозита.\n"
        "На странице: сумма KGS, номер карты получателя, ФИО, банк, кнопки копирования, таймер.\n"
        "3-D Secure нет (P2P-реквизиты, не эквайринг).\n"
        "Вывод: H2H POST /withdrawal/ — account_number = номер карты получателя (16 цифр).\n"
        f"{VIDEO_NOTE}",
        "номер карты получателя, 16 цифр",
    )
    set_comment(
        doc.tables[1].rows[16],
        "Депозит: нет (данные карты отправителя не приходят). Вывод: да — карта в account_number.",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[1].rows[17],
        "Для депозита можем отдать masked номер карты получателя в GET status "
        "(destination_account_number). PAN плательщика не получаем.",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[1].rows[18],
        "non-3DS (P2P по реквизитам, не карточный эквайринг).",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[1].rows[20],
        "Нет (P2P по выданным реквизитам / указанной карте вывода).",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[1].rows[21],
        "Номер карты 16 цифр, держатель, банк (локальные карты KG).",
        prefer_last_n=2,
    )
    set_comment(
        doc.tables[0].rows[28],
        "Депозит: 1) POST /deposit/ method=card2card_kgs → 2) redirect_url → "
        "3) перевод на карту → 4) callback {\"order_id\":\"...\"} → 5) GET status.\n"
        "Вывод: POST /withdrawal/ (account_number=карта) → transaction_id → "
        "callback {\"order_id\"} → GET status.",
    )
    fill_chargeback(doc)
    set_comment(doc.tables[2].rows[4], "Нет (не мобильный метод).")
    set_comment(
        doc.tables[2].rows[5],
        "Masked номер карты получателя можем отдать в GET status (destination_account_number). "
        "PAN плательщика не получаем и не отдаём.",
    )
    fill_fx(doc)
    doc.save(path)


def main() -> None:
    files = {
        "AvaPay_Integration_information_qr_kgs_KG_e497.docx": fill_qr,
        "AvaPay_Integration_information_phone_kgs_KG_924d.docx": fill_phone,
        "AvaPay_Integration_information_card2card_kgs_KG_8c48.docx": fill_card,
    }
    for name, fn in files.items():
        path = DIR / name
        fn(path)
        print(f"updated: {path.name}")


if __name__ == "__main__":
    main()
